"""Table loaders for CSV, TSV, and XLSX inputs.

Design goals (from the master prompt, Phase 2):

* infer delimiters for delimited text,
* preserve sample-ID-like columns as strings (never silently coerce to int),
* detect numeric columns,
* report missing values,
* surface clear validation messages instead of failing silently.
"""

from __future__ import annotations

import csv
import io
import os
from dataclasses import dataclass, field
from typing import Any, BinaryIO, Dict, List, Optional, Sequence, Union

import pandas as pd


class LoaderError(Exception):
    """Raised when a table cannot be loaded or is structurally invalid."""


@dataclass
class TableInfo:
    """A loaded table plus the metadata the app needs to validate mappings.

    For a table read from a specific worksheet of a multi-sheet Excel workbook the
    ``source_*`` provenance fields record where the data came from so that specs,
    exports, and output filenames can stay worksheet-aware. They are ``None`` for
    CSV/TSV/single-table sources (backward compatible).
    """

    dataframe: pd.DataFrame
    source_name: str
    delimiter: Optional[str]
    numeric_columns: List[str] = field(default_factory=list)
    categorical_columns: List[str] = field(default_factory=list)
    missing_value_counts: Dict[str, int] = field(default_factory=dict)
    warnings: List[str] = field(default_factory=list)
    # Worksheet provenance (populated by the workbook loader; None otherwise).
    source_workbook_name: Optional[str] = None
    source_workbook_hash: Optional[str] = None
    source_sheet_name: Optional[str] = None
    source_sheet_index: Optional[int] = None
    source_sheet_type: Optional[str] = None
    source_header_row: Optional[int] = None

    @property
    def columns(self) -> List[str]:
        return list(self.dataframe.columns)

    @property
    def n_rows(self) -> int:
        return int(len(self.dataframe))

    def provenance(self) -> Dict[str, Any]:
        """Worksheet-source provenance for embedding in specs/sidecars.

        Only non-null fields are included so CSV/single-table sources add nothing.
        """
        prov = {
            "source_workbook_name": self.source_workbook_name,
            "source_workbook_hash": self.source_workbook_hash,
            "source_sheet_name": self.source_sheet_name,
            "source_sheet_index": self.source_sheet_index,
            "source_sheet_type": self.source_sheet_type,
            "source_header_row": self.source_header_row,
        }
        return {k: v for k, v in prov.items() if v is not None}

    def summary(self) -> Dict[str, Any]:
        out = {
            "source_name": self.source_name,
            "delimiter": self.delimiter,
            "n_rows": self.n_rows,
            "n_columns": len(self.columns),
            "columns": self.columns,
            "numeric_columns": self.numeric_columns,
            "categorical_columns": self.categorical_columns,
            "missing_value_counts": self.missing_value_counts,
            "warnings": self.warnings,
        }
        prov = self.provenance()
        if prov:
            out["provenance"] = prov
        return out


# Column names that look like identifiers should stay textual even when their
# values happen to be all digits, so that "001" never becomes 1.
_ID_HINTS = ("id", "sample", "patient", "barcode", "accession", "replicate")

_DELIMITER_CANDIDATES = [",", "\t", ";", "|"]


def _looks_like_id_column(name: str) -> bool:
    lowered = str(name).lower()
    return any(hint in lowered for hint in _ID_HINTS)


def _sniff_delimiter(sample_text: str, ext: str) -> str:
    """Guess the field delimiter from a text sample.

    Falls back to the extension hint (.tsv -> tab) and finally a comma.
    """
    try:
        dialect = csv.Sniffer().sniff(sample_text, delimiters="".join(_DELIMITER_CANDIDATES))
        return dialect.delimiter
    except (csv.Error, Exception):  # pragma: no cover - defensive
        pass

    # Heuristic: pick the candidate with the most consistent per-line count.
    lines = [ln for ln in sample_text.splitlines() if ln.strip()][:20]
    best_delim, best_score = None, -1.0
    for cand in _DELIMITER_CANDIDATES:
        counts = [ln.count(cand) for ln in lines]
        if not counts or max(counts) == 0:
            continue
        # reward many columns and penalise inconsistency
        consistency = 1.0 if len(set(counts)) == 1 else 0.5
        score = max(counts) * consistency
        if score > best_score:
            best_delim, best_score = cand, score
    if best_delim is not None:
        return best_delim
    return "\t" if ext == ".tsv" else ","


def _classify_columns(df: pd.DataFrame) -> tuple[List[str], List[str]]:
    numeric, categorical = [], []
    for col in df.columns:
        if pd.api.types.is_numeric_dtype(df[col]) and not _looks_like_id_column(col):
            numeric.append(col)
        else:
            categorical.append(col)
    return numeric, categorical


def _post_process(df: pd.DataFrame, source_name: str, delimiter: Optional[str]) -> TableInfo:
    if df.shape[1] == 0:
        raise LoaderError(f"'{source_name}' has no columns; check the delimiter or file content.")
    if df.shape[0] == 0:
        raise LoaderError(f"'{source_name}' has no data rows.")

    warnings: List[str] = []

    # Strip whitespace from string-like column names.
    df.columns = [str(c).strip() for c in df.columns]

    # Detect duplicate column names, which break mapping by name.
    seen, dups = set(), set()
    for c in df.columns:
        if c in seen:
            dups.add(c)
        seen.add(c)
    if dups:
        warnings.append(f"Duplicate column names found: {sorted(dups)}")

    numeric, categorical = _classify_columns(df)

    missing = {c: int(df[c].isna().sum()) for c in df.columns if df[c].isna().any()}
    if missing:
        warnings.append(
            "Missing values detected in: "
            + ", ".join(f"{c} ({n})" for c, n in missing.items())
        )

    return TableInfo(
        dataframe=df,
        source_name=source_name,
        delimiter=delimiter,
        numeric_columns=numeric,
        categorical_columns=categorical,
        missing_value_counts=missing,
        warnings=warnings,
    )


def table_info_from_dataframe(df: pd.DataFrame, source_name: str) -> TableInfo:
    """Build a :class:`TableInfo` from an in-memory DataFrame.

    Used when a table is produced inside the app (e.g. reshaping a wide matrix to
    long, or adding a grouping column) rather than read from disk. Runs the same
    column classification / missing-value analysis as a file load.
    """
    return _post_process(df.copy(), source_name, delimiter=None)


def _id_dtype_overrides(columns: List[str]) -> Dict[str, str]:
    return {c: "string" for c in columns if _looks_like_id_column(c)}


HEADER_FILL_MODES = ("merged", "forward")


def _normalise_header_rows(header: Union[int, Sequence[int], None]) -> Optional[List[int]]:
    """Return a sorted list of header rows when several were requested, else ``None``.

    A single ``int`` or ``None`` is left for pandas to handle directly; only a genuine multi-row
    request needs the extra combining step, so this keeps the common case untouched.
    """
    if header is None or isinstance(header, (int, bool)):
        return None
    rows = sorted({int(r) for r in header})
    if not rows:
        raise LoaderError("header was given as an empty sequence; pass an int or None.")
    if any(r < 0 for r in rows):
        raise LoaderError(f"header rows must be >= 0, got {rows}.")
    return rows if len(rows) > 1 else None


def _combine_header_rows(df: pd.DataFrame, fill: str = "forward") -> pd.DataFrame:
    """Flatten a pandas MultiIndex header into one unambiguous name per column.

    Spreadsheets written for people routinely put a group label on one row and the sub-label on the
    next, spanning the block it covers, and pandas reports the continuation cells as ``Unnamed: n``.

    ``fill="forward"`` carries the last real label of an outer level rightwards across those blanks
    and is the default *here* because this path handles delimited text, which records no merge
    information at all - a stacked header in a CSV can only be read that way. The Excel path instead
    consults the file's own merged ranges and treats forward filling as opt-in; see
    :func:`_compose_names`.
    """
    if not isinstance(df.columns, pd.MultiIndex):
        return df
    if fill not in HEADER_FILL_MODES:
        raise LoaderError(f"header_fill must be one of {HEADER_FILL_MODES}, got {fill!r}.")

    n_levels = df.columns.nlevels
    filled: List[List[str]] = []
    for level in range(n_levels):
        carried = ""
        values: List[str] = []
        for raw in df.columns.get_level_values(level):
            text = "" if raw is None else str(raw).strip()
            if not text or text.lower().startswith("unnamed:") or text.lower() == "nan":
                # a blank continues the span of the label to its left, but only for the outer
                # levels and only when asked: a blank innermost label has nothing to inherit
                text = carried if (fill == "forward" and level < n_levels - 1) else ""
            else:
                carried = text
            values.append(text)
        filled.append(values)

    names: List[str] = []
    for position in range(df.shape[1]):
        parts = [filled[level][position] for level in range(n_levels)]
        parts = [p for p in parts if p]
        # drop an immediate repeat, so a label merged over a single column is not doubled
        deduped: List[str] = []
        for part in parts:
            if not deduped or deduped[-1] != part:
                deduped.append(part)
        names.append(" | ".join(deduped) if deduped else f"column_{position + 1}")

    out = df.copy()
    out.columns = _uniquify(names)
    return out


def _uniquify(names: Sequence[str]) -> List[str]:
    """Make names unique the way pandas does, so repeated headers stay distinct columns."""
    seen: Dict[str, int] = {}
    out: List[str] = []
    for name in names:
        if name in seen:
            seen[name] += 1
            out.append(f"{name}.{seen[name]}")
        else:
            seen[name] = 0
            out.append(name)
    return out


def _read_delimited(
    buffer: Union[str, bytes],
    source_name: str,
    ext: str,
    header: Union[int, Sequence[int], None] = 0,
    header_fill: str = "forward",
) -> TableInfo:
    """Read delimited text.

    ``header`` is the 0-based row (or rows) holding the column names, or ``None`` for a file with no
    header at all, in which case columns are named positionally so mapping still works. Several
    rows may be given for a stacked header; they are combined by :func:`_combine_header_rows`.
    """
    if isinstance(buffer, bytes):
        text = buffer.decode("utf-8-sig")
    else:
        text = buffer
    if not text.strip():
        raise LoaderError(f"'{source_name}' is empty.")

    sample = "\n".join(text.splitlines()[:50])
    delimiter = _sniff_delimiter(sample, ext)

    multi = _normalise_header_rows(header)
    read_header: Any = multi if multi is not None else header

    # Two-pass read: first read headers to know which columns are IDs, then
    # re-read forcing those columns to stay strings.
    try:
        header_df = pd.read_csv(io.StringIO(text), sep=delimiter, nrows=0,
                                header=read_header, engine="python")
        dtype = _id_dtype_overrides([str(c) for c in header_df.columns])
        df = pd.read_csv(io.StringIO(text), sep=delimiter, dtype=dtype,
                         header=read_header, engine="python")
    except Exception as exc:  # pragma: no cover - defensive
        raise LoaderError(f"Failed to parse '{source_name}' with delimiter {delimiter!r}: {exc}") from exc

    if multi is not None:
        df = _combine_header_rows(df, header_fill)
    elif header is None:
        df.columns = [f"column_{i + 1}" for i in range(df.shape[1])]

    return _post_process(df, source_name, delimiter)


def _merged_header_grid(
    source: Union[str, BinaryIO],
    sheet_name: Optional[Union[str, int]],
    header_rows: Sequence[int],
) -> Optional[List[List[str]]]:
    """Label grid for ``header_rows`` with horizontally merged cells expanded.

    Excel stores a merged range once, in its top-left cell, and leaves the other cells empty, so
    pandas reports every column after the first in a merged block as ``Unnamed: n``. That is why a
    sheet with a genotype label merged across nine replicate columns keeps the label on only the
    first of them.

    The merge ranges are recorded in the file, so the span each label covers is read back exactly
    rather than inferred by filling blanks - which matters because a blank header cell that is *not*
    part of a merge genuinely means "no name", and guessing would silently invent one.

    Returns ``None`` when the merge information cannot be read (a legacy ``.xls``, a stream that
    cannot be rewound, or no openpyxl), in which case the caller keeps pandas' own behaviour.
    """
    try:
        import openpyxl
    except Exception:  # pragma: no cover - openpyxl is a hard dependency in practice
        return None
    try:
        if hasattr(source, "seek"):
            source.seek(0)
        book = openpyxl.load_workbook(source, data_only=True, read_only=False)
    except Exception:
        return None
    try:
        if isinstance(sheet_name, int):
            ws = book.worksheets[sheet_name]
        elif sheet_name is None:
            ws = book.worksheets[0]
        elif sheet_name in book.sheetnames:
            ws = book[sheet_name]
        else:
            return None

        # value of every merged range, keyed by the cells it covers
        spread: Dict[tuple, Any] = {}
        for rng in ws.merged_cells.ranges:
            anchor = ws.cell(row=rng.min_row, column=rng.min_col).value
            if anchor is None:
                continue
            for row in range(rng.min_row, rng.max_row + 1):
                for col in range(rng.min_col, rng.max_col + 1):
                    spread[(row, col)] = anchor

        width = ws.max_column
        grid: List[List[str]] = []
        for header_row in header_rows:
            excel_row = header_row + 1          # openpyxl rows are 1-based
            if excel_row > ws.max_row:
                return None
            labels: List[str] = []
            for col in range(1, width + 1):
                value = spread.get((excel_row, col), ws.cell(row=excel_row, column=col).value)
                labels.append("" if value is None else str(value).strip())
            grid.append(labels)
        return grid
    finally:
        book.close()


def _compose_names(grid: Sequence[Sequence[str]], fill: str = "merged") -> List[str]:
    """Join a label grid into one unambiguous name per column.

    ``fill='merged'`` (the default) trusts only what the file states: a label spans the columns its
    merged range covers, and a blank header cell outside a merge stays blank. That is exact, but a
    spreadsheet author may convey the same grouping by typing the label once and simply leaving the
    neighbouring cells empty, with no merge at all - visually identical, structurally different.

    ``fill='forward'`` additionally carries the last non-blank label rightwards across blanks in the
    outer header rows, which is what a human reads off the sheet. It is opt-in because a blank header
    can equally mean "this column has no name", and inferring a span that the file never recorded
    would be inventing structure.
    """
    if fill not in HEADER_FILL_MODES:
        raise LoaderError(f"header_fill must be one of {HEADER_FILL_MODES}, got {fill!r}.")
    grid = [list(level) for level in grid]
    if fill == "forward" and len(grid) > 1:
        # outer levels only: the innermost label has nothing to inherit
        for level in grid[:-1]:
            carried = ""
            for i, value in enumerate(level):
                text = str(value).strip()
                if text and not text.lower().startswith("unnamed:") and text.lower() != "nan":
                    carried = text
                else:
                    level[i] = carried

    width = len(grid[0]) if grid else 0
    names: List[str] = []
    for position in range(width):
        parts = [str(level[position]).strip() for level in grid]
        parts = [p for p in parts if p and not p.lower().startswith("unnamed:") and p.lower() != "nan"]
        deduped: List[str] = []
        for part in parts:
            if not deduped or deduped[-1] != part:
                deduped.append(part)
        names.append(" | ".join(deduped) if deduped else f"column_{position + 1}")
    return _uniquify(names)


def _read_excel(
    source: Union[str, BinaryIO],
    source_name: str,
    sheet_name: Optional[Union[str, int]],
    header: Union[int, Sequence[int], None] = 0,
    header_fill: str = "merged",
) -> TableInfo:
    """Read one worksheet.

    ``header`` is the 0-based row holding the column names, ``None`` for a sheet with no header, or
    several rows for a stacked header (see :func:`_combine_header_rows`).
    """
    sel = sheet_name if sheet_name is not None else 0
    multi = _normalise_header_rows(header)
    read_header: Any = multi if multi is not None else header

    # Header labels taken from the file's own merged ranges, when they can be read. A single header
    # row merged across a block of replicate columns is the common case and pandas cannot recover it,
    # so this is done for one header row as well as for several.
    header_rows: Optional[List[int]] = None
    if multi is not None:
        header_rows = list(multi)
    elif isinstance(header, int) and not isinstance(header, bool):
        header_rows = [int(header)]
    merged_names: Optional[List[str]] = None
    if header_rows is not None:
        grid = _merged_header_grid(source, sel, header_rows)
        if grid:
            merged_names = _compose_names(grid, header_fill)

    try:
        if merged_names is not None:
            # Names come from the merge-expanded grid, so the data is read positionally from the
            # row after the last header row and the names applied afterwards.
            if hasattr(source, "seek"):
                source.seek(0)
            df = pd.read_excel(source, sheet_name=sel, header=None,
                               skiprows=max(header_rows) + 1)
            if df.shape[1] != len(merged_names):
                merged_names = None          # shape disagreement: fall back to pandas
            else:
                df.columns = merged_names
                dtype = _id_dtype_overrides(merged_names)
                if dtype:
                    for column, kind in dtype.items():
                        if column in df.columns:
                            df[column] = df[column].astype(kind)
        if merged_names is None:
            # Read once to discover columns, then re-read forcing ID columns to str.
            if hasattr(source, "seek"):
                source.seek(0)
            if read_header is None:
                df = pd.read_excel(source, sheet_name=sel, header=None)
            else:
                probe = pd.read_excel(source, sheet_name=sel, header=read_header, nrows=0)
                if hasattr(source, "seek"):
                    source.seek(0)
                dtype = _id_dtype_overrides([str(c) for c in probe.columns])
                df = pd.read_excel(source, sheet_name=sel, header=read_header, dtype=dtype)
    except Exception as exc:
        raise LoaderError(f"Failed to read Excel file '{source_name}': {exc}") from exc

    if isinstance(df, dict):  # multiple sheets returned
        raise LoaderError(
            f"'{source_name}' returned multiple sheets; pass sheet_name to choose one of {list(df.keys())}."
        )
    if merged_names is None and multi is not None:
        df = _combine_header_rows(df, header_fill)
    elif merged_names is None and header is None:
        # No-header read: give columns stable positional names so mapping works.
        df.columns = [f"column_{i + 1}" for i in range(df.shape[1])]
    return _post_process(df, source_name, delimiter=None)


def load_table(
    source: Union[str, os.PathLike, bytes, BinaryIO],
    *,
    source_name: Optional[str] = None,
    file_type: Optional[str] = None,
    sheet_name: Optional[Union[str, int]] = None,
    header: Union[int, Sequence[int], None] = 0,
    header_fill: str = "merged",
) -> TableInfo:
    """Load a CSV/TSV/XLSX table into a :class:`TableInfo`.

    Parameters
    ----------
    source:
        A filesystem path, raw bytes, or a binary file-like object (e.g. a
        Streamlit ``UploadedFile``).
    source_name:
        Display name; inferred from a path when not given.
    file_type:
        One of ``csv``/``tsv``/``xlsx``/``xls``. Inferred from the extension
        when not given. Required when passing raw bytes without a name.
    sheet_name:
        Excel sheet selector (defaults to the first sheet).
    header:
        0-based index of the row holding the column names (the default, ``0``, is the first row),
        ``None`` when the file has no header row, or a sequence of row indices for a stacked
        header. Applies to both Excel and delimited text.
    header_fill:
        How a blank header cell is read for Excel: ``"merged"`` (default) expands only the ranges
        the file records as merged, ``"forward"`` also carries an outer label rightwards across
        blanks. See :func:`_compose_names`.
    """
    # Resolve a name and extension.
    if isinstance(source, (str, os.PathLike)):
        path = os.fspath(source)
        source_name = source_name or os.path.basename(path)
        ext = os.path.splitext(path)[1].lower()
    else:
        name_for_ext = source_name or getattr(source, "name", "") or ""
        source_name = source_name or (name_for_ext if name_for_ext else "uploaded_table")
        ext = os.path.splitext(name_for_ext)[1].lower()

    if file_type:
        ext = "." + file_type.lower().lstrip(".")

    excel_exts = {".xlsx", ".xls", ".xlsm"}
    delimited_exts = {".csv", ".tsv", ".txt", ".tab"}

    # Delimited text records no merged ranges, so a stacked header there can only be read by
    # carrying outer labels forward. "merged" is meaningless for that format and is treated as the
    # forward reading rather than silently producing unnamed columns.
    delim_fill = "forward" if header_fill == "merged" else header_fill

    if ext in excel_exts:
        if isinstance(source, (str, os.PathLike)):
            return _read_excel(os.fspath(source), source_name, sheet_name, header, header_fill)
        if isinstance(source, bytes):
            return _read_excel(io.BytesIO(source), source_name, sheet_name, header, header_fill)
        return _read_excel(source, source_name, sheet_name, header, header_fill)

    if ext not in delimited_exts and ext not in excel_exts:
        # Unknown extension: attempt delimited parsing but warn.
        if ext:
            note = f"Unrecognised extension '{ext}'; attempting delimited text parsing."
        else:
            note = "No file extension; attempting delimited text parsing."
        if isinstance(source, (str, os.PathLike)):
            with open(source, "rb") as fh:
                info = _read_delimited(fh.read(), source_name, ext or ".csv", header, delim_fill)
        elif isinstance(source, bytes):
            info = _read_delimited(source, source_name, ext or ".csv", header, delim_fill)
        else:
            info = _read_delimited(source.read(), source_name, ext or ".csv", header, delim_fill)
        info.warnings.insert(0, note)
        return info

    # Delimited text path.
    if isinstance(source, (str, os.PathLike)):
        with open(source, "rb") as fh:
            return _read_delimited(fh.read(), source_name, ext, header, delim_fill)
    if isinstance(source, bytes):
        return _read_delimited(source, source_name, ext, header, delim_fill)
    return _read_delimited(source.read(), source_name, ext, header, delim_fill)
