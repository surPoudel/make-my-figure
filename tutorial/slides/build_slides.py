"""Build the Make My Figure tutorial / seminar deck from the author's PowerPoint template.

    python tutorial/slides/build_slides.py --template "<path to Suresh_Research_Progress_031826.pptx>"

Rules (from presentation/audit/SLIDE_VISUAL_REVIEW.md): one large visual and one short message per
slide; plots rendered for presentation (tutorial/showcase/*_pres.png: same data, larger typography)
occupy 50-80 % of the usable area; UI screenshots only as readable crops of the control being
taught, or full windows on navigation slides; Arial, titles 28 pt, body 18-20 pt, dark text, one
accent colour. Every placed image is logged to presentation/audit/placements.json for the font QC.
"""
from __future__ import annotations

import argparse
import csv
import json
import os
import sys

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
TUT = os.path.abspath(os.path.join(HERE, ".."))
ROOT = os.path.abspath(os.path.join(TUT, ".."))
SHOTS = os.path.join(TUT, "screenshots")
SHOW = os.path.join(TUT, "showcase")
OUTS = os.path.join(TUT, "automation", "_outputs")
PRES_AUDIT = os.path.join(ROOT, "presentation", "audit")
sys.path.insert(0, ROOT)

FONT = "Arial"
INK = RGBColor(0x11, 0x11, 0x11)
INK2 = RGBColor(0x1F, 0x29, 0x37)
BLUE = RGBColor(0x1D, 0x4E, 0x89)
GREY = RGBColor(0x66, 0x66, 0x66)
RED = RGBColor(0xC0, 0x00, 0x00)
GREEN = RGBColor(0x1B, 0x7A, 0x3E)
LIGHT = RGBColor(0xF2, 0xF4, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

Y_OFF = Inches(0.3)
TITLE_LEFT = Inches(1.85)
CONTENT_TOP = Inches(1.75)      # + Y_OFF -> 2.05 in, below the template header band
placements = []


# ----------------------------------------------------------------------------- helpers
def _clear_slides(prs):
    lst = prs.slides._sldIdLst
    for sid in list(lst):
        prs.part.drop_rel(sid.rId)
        lst.remove(sid)


def _style_runs(tf, size, color=INK, bold=None):
    for p in tf.paragraphs:
        for r in p.runs:
            r.font.name = FONT; r.font.size = Pt(size); r.font.color.rgb = color
            if bold is not None:
                r.font.bold = bold


def add_text(slide, left, top, width, height, text, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, raw=False):
    box = slide.shapes.add_textbox(left, top if raw else top + Y_OFF, width, height)
    tf = box.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    for i, line in enumerate(text if isinstance(text, list) else [text]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = 1.1
        r = p.add_run(); r.text = line
        r.font.name = FONT; r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold
    return box


def add_bullets(slide, left, top, width, height, items, size=18, gap=6):
    box = slide.shapes.add_textbox(left, top + Y_OFF, width, height)
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    for i, it in enumerate(items):
        text, level, bold = (it, 0, False) if isinstance(it, str) else (it + (False,))[:3]
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap); p.line_spacing = 1.08
        if level:
            p.level = 1
        r = p.add_run(); r.text = ("• " if level == 0 else "– ") + text
        r.font.name = FONT; r.font.size = Pt(size - 2 * level); r.font.color.rgb = INK if level == 0 else INK2; r.font.bold = bold
    return box


def add_title(slide, text, size=28):
    add_text(slide, TITLE_LEFT, Inches(0.42), Inches(9.7), Inches(1.1), text, size=size, color=INK2, bold=True,
             anchor=MSO_ANCHOR.MIDDLE, raw=True)


def add_message(slide, text, top=Inches(7.55), size=20, color=INK2):
    add_text(slide, Inches(0.6), top, Inches(10.8), Inches(0.8), text, size=size, color=color, bold=True)


def add_picture_fit(slide, path, left, top, max_w, max_h, crop=None, border=True, kind="ui"):
    img = Image.open(path)
    if crop:
        img = img.crop(crop)
        tmp = os.path.join(HERE, "_crops"); os.makedirs(tmp, exist_ok=True)
        path = os.path.join(tmp, os.path.basename(os.path.dirname(path)) + "_" + os.path.basename(path).replace(".png", f"_{crop[0]}_{crop[1]}_{crop[2]}.png"))
        img.save(path)
    w, h = img.size
    scale = min(max_w / w, max_h / h)
    pw, ph = int(w * scale), int(h * scale)
    l = left + (max_w - pw) // 2; t = top + (max_h - ph) // 2
    pic = slide.shapes.add_picture(path, l, t + Y_OFF, pw, ph)
    if border:
        pic.line.color.rgb = RGBColor(0xC8, 0xC8, 0xC8); pic.line.width = Pt(0.75)
    placements.append({"file": os.path.relpath(path, ROOT), "kind": kind, "placed_width_in": round(pw / 914400, 2),
                       "placed_height_in": round(ph / 914400, 2), "pixel_width": w, "pixel_height": h,
                       "dpi": (img.info.get("dpi") or (None,))[0]})
    return pic


def add_caption(slide, left, top, width, text, size=13):
    add_text(slide, left, top, width, Inches(0.5), text, size=size, color=GREY)


def add_footer(slide, n):
    add_text(slide, Inches(0.5), Inches(8.55), Inches(8), Inches(0.3), f"Make My Figure - desktop tutorial   |   {n}", size=10, color=GREY, raw=True)


def shot(tid, name):
    p = os.path.join(SHOTS, tid, name)
    if not os.path.exists(p):
        raise SystemExit(f"missing screenshot {p}")
    return p


def show(folder, name):
    p = os.path.join(SHOW, folder, name)
    if not os.path.exists(p):
        raise SystemExit(f"missing showcase render {p} - run tutorial/showcase/build_showcase.py")
    return p


def table_excerpt_image(csv_path, rows=6, out_name="table_excerpt.png"):
    """Render the first rows of a CSV as a large, readable table image (data, not UI)."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    with open(csv_path, encoding="utf-8") as fh:
        data = list(csv.reader(fh))[: rows + 1]
    widths = [max(len(str(r[c])) for r in data) + 2 for c in range(len(data[0]))]
    total = float(sum(widths))
    fig, ax = plt.subplots(figsize=(0.135 * total + 0.5, 0.55 * (rows + 1)))
    ax.axis("off")
    tbl = ax.table(cellText=data[1:], colLabels=data[0], loc="center", cellLoc="left",
                   colWidths=[w / total for w in widths])
    tbl.auto_set_font_size(False); tbl.set_fontsize(14); tbl.scale(1, 1.9)
    for (r, c), cell in tbl.get_celld().items():
        cell.set_edgecolor("#CCCCCC")
        if r == 0:
            cell.set_facecolor("#1D4E89"); cell.get_text().set_color("white"); cell.get_text().set_weight("bold")
        elif r % 2 == 0:
            cell.set_facecolor("#F2F4F7")
    tmp = os.path.join(HERE, "_crops"); os.makedirs(tmp, exist_ok=True)
    out = os.path.join(tmp, out_name)
    fig.savefig(out, dpi=200, bbox_inches="tight"); plt.close(fig)
    return out


# ----------------------------------------------------------------------------- content
def build(template, out):
    prs = Presentation(template)
    _clear_slides(prs)
    content_layout = next((l for l in prs.slide_layouts if l.name == "Title and Content"), prs.slide_layouts[1])
    section_layout = next((l for l in prs.slide_layouts if l.name == "Section Header"), prs.slide_layouts[0])
    manifest = json.load(open(os.path.join(TUT, "tutorial_manifest.json"), encoding="utf-8"))
    version, n_plots = manifest["software_version"], manifest["plot_registry_count"]
    n = 0

    def new_slide(layout=content_layout):
        nonlocal n
        s = prs.slides.add_slide(layout)
        for ph in list(s.placeholders):
            ph._element.getparent().remove(ph._element)
        n += 1
        return s

    L = Inches(0.5)
    H_VIS = Inches(5.3)

    # 1 title
    s = new_slide(section_layout)
    add_text(s, Inches(3.0), Inches(3.15), Inches(8.6), Inches(1.1), "Make My Figure", size=40, color=WHITE, bold=True, raw=True)
    add_text(s, Inches(3.0), Inches(4.25), Inches(8.6), Inches(1.5), "One dataset, many controlled views - publication figures with the researcher in control",
             size=22, color=WHITE, raw=True)
    add_text(s, Inches(3.0), Inches(6.35), Inches(8.6), Inches(1.6), ["Desktop tutorial and demonstration", f"Version {version}  -  {n_plots} plot types", "Suresh Poudel"],
             size=18, color=INK2, raw=True)

    # 2 the problem: raw table -> figure
    s = new_slide(); add_title(s, "A real table rarely arrives ready to plot")
    tbl = table_excerpt_image(os.path.join(TUT, "datasets", "ambiguous_columns.csv"), rows=6)
    add_picture_fit(s, tbl, L, CONTENT_TOP, Inches(5.2), H_VIS, border=False, kind="table")
    add_text(s, Inches(5.75), Inches(4.0), Inches(0.6), Inches(0.8), "→", size=40, color=BLUE, align=PP_ALIGN.CENTER)
    add_picture_fit(s, show("1_group_comparison", "J4_black_edged_filled_pres.png"), Inches(6.4), CONTENT_TOP, Inches(5.1), H_VIS, kind="plot")
    add_message(s, "Column names like col_A and measurement_2 are fine: you tell the application what each column means, and the figure follows.")
    add_footer(s, n)

    # 3 what it is (text only)
    s = new_slide(); add_title(s, "What Make My Figure is")
    add_bullets(s, Inches(0.7), CONTENT_TOP, Inches(10.6), Inches(6.0), [
        f"A desktop application (and a browser version) that draws {n_plots} publication-style plot types from CSV, TSV and Excel tables",
        "Runs entirely on your computer - no upload, no account",
        "Proposes column roles, figures and tests; every proposal can be changed",
        "Statistics on the figure: 19 tests, corrections, brackets, effect sizes, a methods sentence",
        "Individual observations with controllable size, jitter, fill, edge and arrangement",
        "Publication presets previewed on your own data before they are applied (experimental library)",
        "Reproducibility: PlotSpec, StatsSpec, Figure Package (.mmfpackage) with frozen data and checksums",
        "Multi-panel Figure Builder; open source, Python, Matplotlib and Qt",
    ], size=19, gap=9)
    add_footer(s, n)

    # 4 design principle
    s = new_slide(); add_title(s, "The application suggests. The researcher decides.")
    for i, (head, body) in enumerate([
        ("It proposes", "Detects column types, recommends figures with a match score, proposes a column for every role, suggests a test from the design."),
        ("You decide", "Every role is a drop-down of your own columns. Change any of them; the figure follows. Choose the test; the panel runs and documents it."),
        ("It records", "PlotSpec for the specification, StatsSpec for the test, Figure Package for specification plus frozen data. Nothing recomputed behind your back."),
    ]):
        x = Inches(0.6) + i * Inches(3.75)
        box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.9) + Y_OFF, Inches(3.5), Inches(4.2))
        box.fill.solid(); box.fill.fore_color.rgb = LIGHT; box.line.color.rgb = RGBColor(0xD0, 0xD5, 0xDD); box.adjustments[0] = 0.06
        add_text(s, x + Inches(0.2), Inches(2.1), Inches(3.1), Inches(0.6), head, size=22, color=BLUE, bold=True)
        add_text(s, x + Inches(0.2), Inches(2.8), Inches(3.1), Inches(3.2), body, size=17)
    add_message(s, "“The data do not belong to one plot.”", top=Inches(6.6), size=22)
    add_footer(s, n)

    # 5 HERO: one dataset, multiple publication-ready views
    s = new_slide(); add_title(s, "One dataset, multiple publication-ready views")
    third = Inches(3.6)
    for i, (name, cap) in enumerate([("B_box_points_outline_preset_pres.png", "Box + observations (outline)"),
                                     ("C_violin_points_preset_pres.png", "Violin + observations"),
                                     ("D_bar_points_jittered_preset_pres.png", "Bar + observations (jittered)")]):
        x = L + i * (third + Inches(0.1))
        add_picture_fit(s, show("1_group_comparison", name), x, CONTENT_TOP, third, Inches(4.9), border=False, kind="plot")
        add_caption(s, x, Inches(6.75), third, cap, size=14)
    add_message(s, "Same 40 observations (n = 7, 9, 11, 13), same Welch tests and P values - three experimental publication presets.", top=Inches(7.3), size=19)
    add_footer(s, n)

    # 6 observations under your control
    s = new_slide(); add_title(s, "The observations are yours to show - size, jitter, fill, edge, arrangement")
    names = [("J1_small_narrow_jitter_pres.png", "small, narrow jitter"), ("J2_large_moderate_jitter_pres.png", "large, moderate jitter"),
             ("J3_open_circles_pres.png", "open circles"), ("J4_black_edged_filled_pres.png", "filled, dark edge"),
             ("J5_beeswarm_pres.png", "beeswarm"), ("J6_centered_no_jitter_pres.png", "centred, translucent, no jitter")]
    cw, ch = Inches(3.6), Inches(2.75)
    for i, (name, cap) in enumerate(names):
        r, c = divmod(i, 3)
        x = L + c * (cw + Inches(0.1)); y = CONTENT_TOP + r * (ch + Inches(0.35))
        add_picture_fit(s, show("1_group_comparison", name), x, y, cw, ch - Inches(0.05), border=False, kind="plot")
        add_caption(s, x, y + ch - Inches(0.05), cw, cap, size=13)
    add_message(s, "The scientific data do not change. The presentation can.", top=Inches(7.85), size=19)
    add_footer(s, n)

    # 7 statistics
    s = new_slide(); add_title(s, "Statistics on the figure - chosen by you, documented by the application")
    add_picture_fit(s, show("1_group_comparison", "B_box_points_outline_preset_pres.png"), L, CONTENT_TOP, Inches(5.5), Inches(5.4), border=False, kind="plot")
    add_picture_fit(s, shot("group_comparison_box", "05_statistics.png"), Inches(6.2), CONTENT_TOP, Inches(5.4), Inches(5.4), crop=(0, 20, 566, 520), kind="ui")
    add_message(s, "Welch's t-test, selected pairs, Holm correction: brackets with exact P, a results table, and the methods sentence to paste.", top=Inches(7.5), size=18)
    add_footer(s, n)

    # 8a publication presets: the preview, before / after on your data
    s = new_slide(); add_title(s, "Publication presets are previewed on your own data")
    add_picture_fit(s, shot("observations_jitter", "09_preview_dialog.png"), L, CONTENT_TOP, Inches(11.0), Inches(4.9), crop=(0, 0, 908, 410), kind="ui")
    add_message(s, "Before and after, drawn on the loaded table - nothing is applied until you confirm.", top=Inches(7.2), size=19)
    add_footer(s, n)

    # 8b publication presets: what changes, what is protected
    s = new_slide(); add_title(s, "The preview lists what changes and checks what must not")
    add_picture_fit(s, shot("observations_jitter", "09_preview_dialog.png"), L, CONTENT_TOP, Inches(11.0), Inches(2.8), crop=(0, 410, 908, 638), kind="ui")
    add_bullets(s, Inches(0.7), Inches(5.0), Inches(10.6), Inches(2.8), [
        "Twelve experimental presets derived from published open-access figures and publishers' stated requirements; the letter in a name, (N) (S) (C), is the evidence set, not an approval",
        "Every setting that would change is listed (here marker opacity, point edge width, point size)",
        "Green line: no data, column role, statistical test, threshold or transformation changes - a preset that would is refused",
        "Apply, Cancel, or Save as my preset - a preset is a starting configuration, not a locked template",
    ], size=17, gap=8)
    add_footer(s, n)

    # 9-12 same data, different presentation
    pairs = [
        ("Same data, different presentation - volcano", "2_volcano", "A_default_pres.png", "B_refined_pres.png",
         "default: raw P, ten labels", "refined: adjusted P (FDR), boxed labels, larger markers",
         "1,200 features, identical values; the axis, labels and markers changed, the numbers did not."),
        ("Same data, different presentation - clustered heatmap", "3_heatmap", "A_basic_pres.png", "B_refined_pres.png",
         "basic: unscaled values, auto colormap, clustered columns", "refined: row z-score, RdBu_r, samples in file order",
         "The same 30 most variable features. Scaling and colormap are presentation choices the reader must be told about."),
        ("Same data, different presentation - scatter with regression", "4_scatter", "A_basic_pres.png", "B_refined_pres.png",
         "basic: two roles, one fit", "refined: colour by cell line, per-group fit, r / R² / P / n",
         "Adding a colour role turns one regression into one per group; the sixty points are the same."),
        ("Same data, different presentation - Kaplan-Meier", "5_survival", "A_default_pres.png", "B_refined_pres.png",
         "default: fraction, auto ticks", "refined: percent, 50 % reference, log-rank P on the figure",
         "Eighty subjects, the same curves; the log-rank test is written on the figure and stored with it."),
    ]
    for title, folder, a, b, ca, cb, msg in pairs:
        s = new_slide(); add_title(s, title)
        half = Inches(5.4)
        add_picture_fit(s, show(folder, a), L, CONTENT_TOP, half, Inches(4.9), border=False, kind="plot")
        add_picture_fit(s, show(folder, b), L + half + Inches(0.2), CONTENT_TOP, half, Inches(4.9), border=False, kind="plot")
        add_caption(s, L, Inches(6.75), half, ca, size=14); add_caption(s, L + half + Inches(0.2), Inches(6.75), half, cb, size=14)
        add_message(s, msg, top=Inches(7.3), size=18)
        add_footer(s, n)

    # 13 HERO: from raw table to figure
    s = new_slide(); add_title(s, "From an ambiguous table to a publication figure")
    tbl2 = table_excerpt_image(os.path.join(TUT, "datasets", "showcase_group_comparison.csv"), rows=6, out_name="showcase_excerpt.png")
    add_picture_fit(s, tbl2, L, CONTENT_TOP, Inches(5.3), Inches(2.3), border=False, kind="table")
    add_caption(s, L, Inches(4.1), Inches(5.3), "1  open the table: 40 animals, 4 groups, unequal n", size=14)
    add_text(s, Inches(5.85), Inches(2.6), Inches(0.5), Inches(0.8), "\u2192", size=36, color=BLUE, align=PP_ALIGN.CENTER)
    add_picture_fit(s, shot("observations_jitter", "07b_options_after.png"), Inches(6.2), CONTENT_TOP, Inches(5.4), Inches(2.3), crop=(0, 0, 566, 240), kind="ui")
    add_caption(s, Inches(6.2), Inches(4.1), Inches(5.4), "2  map the roles; choose how the observations are drawn", size=14)
    add_picture_fit(s, show("1_group_comparison", "B_box_points_outline_preset_pres.png"), L, Inches(4.5), Inches(3.9), Inches(3.4), border=False, kind="plot")
    add_caption(s, L, Inches(7.85), Inches(4.0), "3  statistics, publication preset, export", size=14)
    add_text(s, Inches(4.9), Inches(5.0), Inches(6.7), Inches(2.6),
             ["Open, map, draw, test, preset, export.", "",
              "Five minutes from a table with unnamed columns to a figure with brackets, a methods sentence, its PlotSpec and a Figure Package."],
             size=19, color=INK2, bold=False)
    add_footer(s, n)

    # 14 volcano manual mapping
    s = new_slide(); add_title(s, "When the detector does not know your headers, you assign the roles")
    add_picture_fit(s, shot("volcano_manual_mapping", "04b_window_before.png"), L, CONTENT_TOP, Inches(11.0), Inches(1.3), crop=(612, 48, 1680, 125), kind="ui")
    add_picture_fit(s, shot("volcano_manual_mapping", "04_mapping_renamed_before.png"), L, Inches(3.3), Inches(5.4), Inches(1.45), kind="ui")
    add_caption(s, L, Inches(4.75), Inches(5.4), "every role starts at (none): the detector does not know these headers", size=13)
    add_picture_fit(s, shot("volcano_manual_mapping", "05_mapping_renamed_after.png"), L, Inches(5.25), Inches(5.4), Inches(1.45), kind="ui")
    add_caption(s, L, Inches(6.7), Inches(5.4), "x, p, label, id_col assigned by hand", size=13)
    add_picture_fit(s, show("2_volcano", "B_refined_pres.png"), Inches(6.1), Inches(3.3), Inches(5.5), Inches(4.1), border=False, kind="plot")
    add_message(s, "The Messages tab names the missing roles; four drop-downs later the volcano renders - P values read from the table, never recomputed.", top=Inches(7.55), size=17)
    add_footer(s, n)

    # 15 reproducibility
    s = new_slide(); add_title(s, "Reproducibility - a Figure Package carries the figure and its frozen data")
    add_picture_fit(s, shot("figure_package", "03_package_confirmation.png"), L, CONTENT_TOP, Inches(5.0), Inches(4.2), kind="ui")
    add_bullets(s, Inches(5.9), CONTENT_TOP, Inches(5.8), Inches(4.4), [
        ("What is inside one .mmfpackage", 0, True),
        ("plot_spec.json, stats_spec.json (settings and results), render_metadata.json", 1),
        ("data/source.csv and the typed table; the original file when available", 1),
        ("preview/figure.png, .svg, .pdf; environment/environment.json", 1),
        ("manifest.json with a SHA-256 for every member - verified before drawing", 1),
        ("PlotSpec alone is the recipe; it needs the data file to redraw", 0, True),
    ], size=16, gap=6)
    add_message(s, "Move the one file anywhere, open it: integrity verified, mapping, options and brackets restored.", top=Inches(6.6), size=19)
    add_footer(s, n)

    # 16 figure builder
    s = new_slide(); add_title(s, "Figure Builder - panels carry their own data")
    comp = os.path.join(OUTS, "figure_builder", "composite.png")
    if os.path.exists(comp):
        add_picture_fit(s, comp, L, CONTENT_TOP, Inches(11.0), Inches(5.2), border=False, kind="plot")
    add_message(s, "Panel letters, width in millimetres, fonts for all panels; the composite exports as PNG, PDF, SVG and as a Figure Package.", top=Inches(7.5), size=18)
    add_footer(s, n)

    # 17 tutorial system
    s = new_slide(); add_title(s, "The tutorial is generated from the real application and checked by running it")
    x = Inches(0.6)
    for i, t in enumerate(["tutorial action script", "driver operates the real MainWindow", "screenshots (offscreen) or video (on screen)", "validation record PASS / FAIL"]):
        box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.6) + Y_OFF, Inches(2.5), Inches(1.2))
        box.fill.solid(); box.fill.fore_color.rgb = LIGHT if i % 2 == 0 else RGBColor(0xE3, 0xEB, 0xF5); box.line.color.rgb = RGBColor(0xC8, 0xD0, 0xDC)
        add_text(s, x + Inches(0.1), Inches(2.7), Inches(2.3), Inches(1.0), t, size=16, color=INK2, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        if i < 3:
            add_text(s, x + Inches(2.5), Inches(2.85), Inches(0.4), Inches(0.6), "→", size=24, color=BLUE, align=PP_ALIGN.CENTER)
        x += Inches(2.85)
    add_message(s, "No mock-ups: every screenshot is a grab of the production widgets; every step is a step a user can click; every claim was executed.", top=Inches(4.6), size=19)
    add_bullets(s, Inches(0.7), Inches(5.6), Inches(10.6), Inches(2.4), [
        "Same script, two modes: deterministic screenshots, or a visible window for video recording (recording kept separate)",
        "Same-data showcases rendered by the application's renderer; presets applied from their files; data integrity recorded per pair",
    ], size=16)
    add_footer(s, n)

    # 18 validation
    s = new_slide(); add_title(s, "Every tutorial instruction executed against the application")
    rows = [("Tutorial", "Status", "Checks")]
    for p in sorted(os.listdir(os.path.join(TUT, "audit", "validation"))):
        r = json.load(open(os.path.join(TUT, "audit", "validation", p), encoding="utf-8"))
        rows.append((r["title"], r["status"], f"{len(r['checks'])} checks, {len(r['captures'])} captures"))
    rh = Inches(0.38)
    tbl = s.shapes.add_table(len(rows), 3, Inches(0.6), CONTENT_TOP + Y_OFF, Inches(10.8), rh * len(rows)).table
    tbl.columns[0].width = Inches(6.2); tbl.columns[1].width = Inches(1.4); tbl.columns[2].width = Inches(3.2)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = tbl.cell(i, j); cell.text = val
            _style_runs(cell.text_frame, 13 if i else 14, INK if i else WHITE, bold=(i == 0))
            cell.fill.solid(); cell.fill.fore_color.rgb = BLUE if i == 0 else (WHITE if i % 2 else LIGHT)
            if j == 1 and i:
                _style_runs(cell.text_frame, 13, GREEN if val == "PASS" else RED, bold=True)
    add_footer(s, n)

    # 19 mosaic
    s = new_slide(); add_title(s, f"{n_plots} plot types, one workflow")
    gal = json.load(open(os.path.join(TUT, "audit", "gallery_index.json"), encoding="utf-8"))
    thumbs = [g for g in gal if g["thumb"]]
    cw, ch = Inches(1.38), Inches(1.28)
    for i, g in enumerate(thumbs[:40]):
        r_, c_ = divmod(i, 8)
        add_picture_fit(s, os.path.join(TUT, g["thumb"]), Inches(0.45) + c_ * cw, Inches(1.65) + r_ * ch, cw - Inches(0.08), ch - Inches(0.08), border=False, kind="thumbnail")
    add_caption(s, Inches(0.5), Inches(8.1), Inches(11), "Thumbnails for breadth only - each is a real render of the bundled synthetic example (tutorial/PLOT_GALLERY.md)")
    add_footer(s, n)

    # 20 next
    s = new_slide(); add_title(s, "Where this goes next")
    add_bullets(s, Inches(0.7), CONTENT_TOP, Inches(10.6), Inches(6.0), [
        ("Reviewed and rebuilt", 0, True),
        ("tutorial and slides re-evaluated as visual demonstrations; observations, jitter and publication presets shown at readable size", 1),
        ("After author review", 0, True),
        ("remaining plot tutorials with the same template; matrix workflow tutorial and video; master and plot videos recorded live", 1),
        ("decision whether the presets branch (observation controls, experimental presets) is merged - without it these chapters cannot ship", 1),
        ("Not merged, tagged, released, published or pushed", 0, True),
    ], size=19, gap=9)
    add_footer(s, n)

    # 21 summary
    s = new_slide(); add_title(s, "Summary")
    add_bullets(s, Inches(0.7), CONTENT_TOP, Inches(10.6), Inches(5.2), [
        "Open any table; the application detects, recommends and proposes - and you assign every role yourself",
        "One dataset, many controlled representations: summary plus every observation, with the test on the figure",
        "Publication presets are previewed on your data and remain a starting point",
        "PlotSpec carries the recipe; a Figure Package carries recipe and ingredients, verified",
        "The tutorial is generated from the real application and validated by running it",
    ], size=19, gap=9)
    add_text(s, Inches(0.7), Inches(6.9), Inches(10.6), Inches(1.0), "github.com/surPoudel/make-my-figure  -  tutorial/ on the pilot branch (not yet published)", size=15, color=GREY)
    add_footer(s, n)

    prs.save(out)
    os.makedirs(PRES_AUDIT, exist_ok=True)
    with open(os.path.join(PRES_AUDIT, "placements.json"), "w", encoding="utf-8") as fh:
        json.dump(placements, fh, indent=1)
    print(f"saved {out} ({n} slides); {len(placements)} placed images logged")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--template", required=True)
    ap.add_argument("--out", default=os.path.join(HERE, "MakeMyFigure_Tutorial_Seminar.pptx"))
    a = ap.parse_args()
    build(a.template, a.out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
